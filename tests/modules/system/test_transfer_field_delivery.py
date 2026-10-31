import asyncio
from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import insert, select, update

from app.core.authorization_plan import LockMode, TenantLock
from app.core.exceptions import AuthorizationException, BusinessRuleException
from app.core.file_storage import LocalFileStorage
from app.db.authorization_transfer_owner import TransferOwner
from app.db.authorization_transfer_read_owner import TransferReadOwner
from app.db.base import user_depts
from app.modules.system.field_policy.evaluator import FieldGrant
from app.modules.system.models.dept import Dept
from app.modules.system.models.role import Role
from app.modules.system.models.user_field_policy import UserFieldPolicyState
from app.modules.system.models.user_transfer import UserExportTask, UserImportBatch
from app.modules.system.service.authorization_transaction_service import (
    authorization_transactions,
)
from app.modules.system.service.user_field_context import user_field_context
from app.modules.system.service.user_transfer_fields import user_transfer_fields
from tests.modules.system.test_authorization_import_candidate import (
    import_sessions as import_sessions,
)


async def test_export_records_sources_and_rejects_epoch_change_before_file_read(
    import_sessions, tmp_path, monkeypatch
):
    factory, plain, tenant = import_sessions
    storage = LocalFileStorage(tmp_path)
    owner = TransferOwner(factory, storage)
    result = await owner.export(
        tenant=tenant, actor_user_version=1, args={"reason": "audit"}
    )
    async with plain() as db:
        row = await db.get(UserExportTask, result.resource_id)
        assert row.field_lineage["export"]["sources"]
        assert "field_lineage" not in row.filter_snapshot
    data, _ = await owner.download(
        tenant=tenant,
        actor_user_version=1,
        kind="export",
        identifier=result.resource_id,
        administrative=True,
    )
    assert data == result.data
    async with plain() as db:
        await db.execute(
            update(UserFieldPolicyState).values(epoch=UserFieldPolicyState.epoch + 1)
        )
        await db.commit()

    async def must_not_read(*_args, **_kwargs):
        pytest.fail("revoked download touched the file store")

    monkeypatch.setattr(storage, "artifact_identity", must_not_read)
    with pytest.raises(AuthorizationException, match="重新生成"):
        await owner.download(
            tenant=tenant,
            actor_user_version=1,
            kind="export",
            identifier=result.resource_id,
            administrative=True,
        )


async def test_import_denial_precedes_preview_insert(
    import_sessions, tmp_path, monkeypatch
):
    factory, plain, tenant = import_sessions

    async def denied_context(db, tx, *, tenant, version):
        return await user_field_context.load(
            db,
            tx,
            tenant=tenant,
            actor_user_version=version,
            grants=(FieldGrant("201", "user_name", "exact", "import"),),
        )

    monkeypatch.setattr(user_transfer_fields, "context", denied_context)
    with pytest.raises(AuthorizationException) as caught:
        await TransferOwner(factory, LocalFileStorage(tmp_path)).preview(
            tenant=tenant,
            actor_user_version=1,
            file_bytes=b"user_name,dept_input\nnewperson,Team\n",
            filename="input.csv",
            reason="audit",
        )
    assert caught.value.error_code == "FIELD_WRITE_DENIED"
    async with plain() as db:
        assert await db.scalar(select(UserImportBatch.batch_id)) is None
    assert not list(tmp_path.rglob("payload"))


async def test_export_actual_workbook_obeys_trusted_denies(
    import_sessions, tmp_path, monkeypatch
):
    factory, plain, tenant = import_sessions

    async def limited_context(db, tx, *, tenant, version):
        return await user_field_context.load(
            db,
            tx,
            tenant=tenant,
            actor_user_version=version,
            grants=(FieldGrant("201", "user_name", "export", "export"),),
        )

    monkeypatch.setattr(user_transfer_fields, "context", limited_context)
    result = await TransferOwner(factory, LocalFileStorage(tmp_path)).export(
        tenant=tenant, actor_user_version=1, args={"reason": "audit"}
    )
    wb = load_workbook(BytesIO(result.data))
    assert "importer" not in str(
        [cell.value for ws in wb for row in ws for cell in row]
    )
    wb.close()


async def test_template_dictionary_excludes_inaccessible_departments(import_sessions):

    factory, plain, tenant = import_sessions
    async with plain() as db:
        await db.execute(
            update(Role)
            .where(Role.role_id == 201)
            .values(role_code="LIMITED", data_scope="2")
        )
        await db.execute(
            insert(user_depts).values(
                tenant_id=0, user_id=101, dept_id=301, is_primary="1"
            )
        )
        await db.execute(
            insert(Dept).values(
                tenant_id=0,
                dept_id=302,
                dept_name="secret-department",
                ancestors="0",
                status="1",
            )
        )
        await db.commit()
    data = await TransferReadOwner(factory).execute(
        tenant=tenant, actor_user_version=1, operation="template"
    )
    wb = load_workbook(BytesIO(data))
    assert len(wb.worksheets) == 4
    assert "secret-department" not in str(
        [cell.value for ws in wb for row in ws for cell in row]
    )
    wb.close()


async def test_download_projection_holds_barrier_until_bytes_are_frozen(
    import_sessions, tmp_path, monkeypatch
):
    factory, _, tenant = import_sessions
    storage = LocalFileStorage(tmp_path)
    owner = TransferOwner(factory, storage)
    result = await owner.export(
        tenant=tenant, actor_user_version=1, args={"reason": "audit"}
    )
    entered, release, writer_started, writer_entered = (
        asyncio.Event() for _ in range(4)
    )
    original = storage.read_artifact

    async def blocked(*args, **kwargs):
        entered.set()
        await release.wait()
        return await original(*args, **kwargs)

    monkeypatch.setattr(storage, "read_artifact", blocked)

    async def writer():
        async with factory() as db:
            writer_started.set()
            async with authorization_transactions.scope(
                db, tenants=(TenantLock(0, LockMode.EXCLUSIVE),)
            ):
                writer_entered.set()
            await db.rollback()

    downloading = asyncio.create_task(
        owner.download(
            tenant=tenant,
            actor_user_version=1,
            kind="export",
            identifier=result.resource_id,
            administrative=True,
        )
    )
    writing = None
    try:
        await asyncio.wait_for(entered.wait(), 5)
        writing = asyncio.create_task(writer())
        await asyncio.wait_for(writer_started.wait(), 5)
        await asyncio.sleep(0.05)
        assert not writer_entered.is_set()
        release.set()
        data, _ = await downloading
        await asyncio.wait_for(writing, 5)
        assert data == result.data and writer_entered.is_set()
    finally:
        release.set()
        await asyncio.gather(
            downloading, *([writing] if writing else []), return_exceptions=True
        )


async def test_preview_source_rechecks_write_fields_even_with_same_epoch(
    import_sessions, tmp_path, monkeypatch
):
    factory, _, tenant = import_sessions
    storage = LocalFileStorage(tmp_path)
    owner = TransferOwner(factory, storage)
    preview = await owner.preview(
        tenant=tenant,
        actor_user_version=1,
        file_bytes=b"user_name,dept_input\nnewperson,Team\n",
        filename="users.csv",
        reason="audit",
    )

    async def denied_context(db, tx, *, tenant, version):
        return await user_field_context.load(
            db, tx, tenant=tenant, actor_user_version=version, grants=()
        )

    monkeypatch.setattr(user_transfer_fields, "context", denied_context)

    async def unopened(*_args, **_kwargs):
        pytest.fail("denied source download opened storage")

    monkeypatch.setattr(storage, "artifact_identity", unopened)
    with pytest.raises(AuthorizationException):
        await owner.download(
            tenant=tenant,
            actor_user_version=1,
            kind="preview_source",
            identifier=preview.resource_id,
        )


async def test_export_filter_is_proved_before_it_can_hide_counterexamples(
    import_sessions, tmp_path, monkeypatch
):
    factory, plain, tenant = import_sessions

    async def denied_context(db, tx, *, tenant, version):
        return await user_field_context.load(
            db, tx, tenant=tenant, actor_user_version=version, grants=()
        )

    monkeypatch.setattr(user_transfer_fields, "context", denied_context)
    with pytest.raises(AuthorizationException) as caught:
        await TransferOwner(factory, LocalFileStorage(tmp_path)).export(
            tenant=tenant,
            actor_user_version=1,
            args={"reason": "audit", "user_email": "no-match@example.com"},
        )
    assert caught.value.error_code == "FIELD_QUERY_DENIED"
    async with plain() as db:
        assert await db.scalar(select(UserExportTask.export_id)) is None
    assert not list(tmp_path.rglob("payload"))


async def test_employee_match_requires_import_exact_before_preview(
    import_sessions, tmp_path, monkeypatch
):
    factory, plain, tenant = import_sessions

    async def denied_context(db, tx, *, tenant, version):
        return await user_field_context.load(
            db,
            tx,
            tenant=tenant,
            actor_user_version=version,
            grants=(FieldGrant("201", "user_name", "exact", "import"),),
        )

    monkeypatch.setattr(user_transfer_fields, "context", denied_context)
    with pytest.raises(AuthorizationException) as caught:
        await TransferOwner(factory, LocalFileStorage(tmp_path)).preview(
            tenant=tenant,
            actor_user_version=1,
            file_bytes=b"user_name,dept_input,employee_no\nnewperson,Team,nonexistent\n",
            filename="users.csv",
            reason="audit",
        )
    assert caught.value.error_code == "FIELD_QUERY_DENIED"
    async with plain() as db:
        assert await db.scalar(select(UserImportBatch.batch_id)) is None


async def test_missing_lineage_only_has_explicit_legacy_compatibility(
    import_sessions, tmp_path, monkeypatch
):
    factory, plain, tenant = import_sessions
    storage = LocalFileStorage(tmp_path)
    owner = TransferOwner(factory, storage)
    result = await owner.export(
        tenant=tenant, actor_user_version=1, args={"reason": "audit"}
    )
    async with plain() as db:
        await db.execute(update(UserExportTask).values(field_lineage=None))
        await db.commit()
    data, _ = await owner.download(
        tenant=tenant,
        actor_user_version=1,
        kind="export",
        identifier=result.resource_id,
        administrative=True,
    )
    assert data == result.data
    monkeypatch.setattr(user_transfer_fields, "legacy_assets", False)
    with pytest.raises(AuthorizationException):
        await owner.download(
            tenant=tenant,
            actor_user_version=1,
            kind="export",
            identifier=result.resource_id,
            administrative=True,
        )


async def test_batch_metadata_uses_current_operator_field_and_hides_internal_lineage(
    import_sessions, tmp_path, monkeypatch
):
    factory, _, tenant = import_sessions
    preview = await TransferOwner(factory, LocalFileStorage(tmp_path)).preview(
        tenant=tenant,
        actor_user_version=1,
        file_bytes=b"user_name,dept_input\nnewperson,Team\n",
        filename="secret-filename.csv",
        reason="secret-reason",
    )
    owner = TransferReadOwner(factory)
    result = await owner.execute(
        tenant=tenant,
        actor_user_version=1,
        operation="batch",
        identifier=preview.resource_id,
    )
    assert result["operatorName"] == "importer"
    assert (
        "secret" not in str(result)
        and "fieldLineage" not in result
        and "field_lineage" not in result
    )

    async def hidden(db, tx, *, tenant, version):
        return await user_field_context.load(
            db, tx, tenant=tenant, actor_user_version=version, grants=()
        )

    monkeypatch.setattr(user_transfer_fields, "context", hidden)
    result = await owner.execute(
        tenant=tenant,
        actor_user_version=1,
        operation="batch",
        identifier=preview.resource_id,
    )
    assert result["operatorName"] is None


async def test_rest_download_checks_retention_before_storage(
    import_sessions, tmp_path, monkeypatch
):
    factory, plain, tenant = import_sessions
    storage = LocalFileStorage(tmp_path)
    owner = TransferOwner(factory, storage)
    result = await owner.export(
        tenant=tenant, actor_user_version=1, args={"reason": "audit"}
    )
    async with plain() as db:
        await db.execute(
            update(UserExportTask).values(
                created_at=(datetime.now(UTC) - timedelta(days=31)).replace(tzinfo=None)
            )
        )
        await db.commit()

    async def unopened(*_args, **_kwargs):
        pytest.fail("expired download opened storage")

    monkeypatch.setattr(storage, "artifact_identity", unopened)
    with pytest.raises(BusinessRuleException) as caught:
        await owner.download(
            tenant=tenant,
            actor_user_version=1,
            kind="export",
            identifier=result.resource_id,
            administrative=True,
        )
    assert caught.value.error_code == "AI_EXPORT_FILE_EXPIRED"
