"""Sensitive config values: masked in list/export; mask sentinel keeps the value on update."""

import io

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession
from tenant_helpers import create_test_tenant, tenant_context

from app.modules.system.models.config import Config
from app.modules.system.schemas.config import ConfigOut, ConfigQuery, ConfigUpdate
from app.modules.system.service.config_service import (
    MASKED_CONFIG_VALUE,
    config_service,
)


async def _setup(db_session: AsyncSession, *, keys: list[tuple[str, str]]):
    tenant_row = await create_test_tenant(db_session, prefix="cfg-mask")
    tenant = tenant_context(tenant_id=tenant_row.tenant_id, actor_user_id=1)
    configs = []
    for key, value in keys:
        config = Config(
            tenant_id=tenant.tenant_id,
            config_name=key,
            config_key=key,
            config_value=value,
            config_type="text",
            config_group="auth" if key.startswith("auth:") else "feature",
            status="1",
            is_public=False,
        )
        db_session.add(config)
        configs.append(config)
    await db_session.flush()
    return tenant, configs


async def test_list_masks_sensitive_values_only(db_session: AsyncSession) -> None:
    tenant, _ = await _setup(
        db_session,
        keys=[("auth:default_password", "Hohu123456"), ("feature:toggle", "true")],
    )

    page = await config_service.get_list(
        db_session, ConfigQuery(size=50), tenant=tenant
    )
    # Serialize via ConfigOut like the API does (catches missing required fields).
    outs = [ConfigOut.model_validate(c) for c in page.records]
    by_key = {o.config_key: o.config_value for o in outs}

    assert by_key["auth:default_password"] == MASKED_CONFIG_VALUE
    assert by_key["feature:toggle"] == "true"


async def test_update_with_mask_sentinel_keeps_original_value(
    db_session: AsyncSession,
) -> None:
    tenant, configs = await _setup(
        db_session, keys=[("auth:default_password", "S3cret!")]
    )
    config = configs[0]

    await config_service.update(
        db_session,
        config.config_id,
        ConfigUpdate(config_value=MASKED_CONFIG_VALUE),
        tenant=tenant,
    )
    await db_session.flush()
    await db_session.refresh(config)
    assert config.config_value == "S3cret!"

    await config_service.update(
        db_session,
        config.config_id,
        ConfigUpdate(config_value="NewPass@123"),
        tenant=tenant,
    )
    await db_session.flush()
    await db_session.refresh(config)
    assert config.config_value == "NewPass@123"


async def test_export_masks_sensitive_values(db_session: AsyncSession) -> None:
    tenant, _ = await _setup(
        db_session,
        keys=[("auth:default_password", "Hohu123456"), ("feature:toggle", "true")],
    )

    buf = await config_service.export_configs(
        db_session, ConfigQuery(size=100), tenant=tenant
    )
    rows = {
        row[1]: row[2]
        for row in load_workbook(io.BytesIO(buf.getvalue())).active.iter_rows(
            values_only=True
        )
        if row[1] not in (None, "config_key")
    }
    assert rows["auth:default_password"] == MASKED_CONFIG_VALUE
    assert rows["feature:toggle"] == "true"
