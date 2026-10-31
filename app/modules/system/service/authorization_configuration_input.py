"""Bounded, immutable inputs for dormant configuration candidates."""

import json
from dataclasses import dataclass, field
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, ValidationError

from app.core.authorization_io import BoundedWorker
from app.core.exceptions import BusinessException, BusinessRuleException
from app.modules.system.models.config import Config
from app.modules.system.models.menu import Menu
from app.modules.system.schemas.config import ConfigCreate, ConfigUpdate
from app.modules.system.schemas.menu import MenuCreate, MenuUpdate
from app.modules.system.service.mail_service import MailWrite
from app.modules.system.service.settings_service import _validate
from app.modules.system.settings_catalog import SETTING_GROUPS, SETTINGS
from app.utils.safe_xlsx import UnsafeXlsxError, validate_untrusted_xlsx_xml

_IMPORT_ISSUED = object()
_import_worker = BoundedWorker(workers=4, queue=32, timeout=3.0)


def invalid(code="CONFIGURATION_INPUT_INVALID"):
    return BusinessRuleException("Invalid configuration input", error_code=code)


def limit():
    return BusinessException(
        422,
        "Configuration exceeds the input budget",
        error_code="AUTHORIZATION_PLAN_LIMIT_EXCEEDED",
    )


def identifier(value):
    if type(value) is str and value.isascii() and value.isdecimal():
        value = int(value)
    if type(value) is not int or not 0 < value < 2**63:
        raise invalid()
    return value


def bounded_json(value, maximum):
    # Reject arbitrary objects, cycles and deeply nested input before encoding.
    nodes = 0

    def visit(item, depth=0):
        nonlocal nodes
        nodes += 1
        if depth > 8 or nodes > 16384:
            raise limit()
        if type(item) is dict:
            if any(type(key) is not str for key in item):
                raise invalid()
            for child in item.values():
                visit(child, depth + 1)
        elif type(item) in {list, tuple}:
            for child in item:
                visit(child, depth + 1)
        elif type(item) not in {str, bool, int, float, type(None)}:
            raise invalid()
        elif type(item) is str and len(item.encode("utf-8")) > maximum:
            raise limit()

    visit(value)
    try:
        encoded = json.dumps(value, ensure_ascii=False, allow_nan=False)
    except (TypeError, ValueError):
        raise invalid() from None
    if len(encoded.encode("utf-8")) > maximum:
        raise limit()
    return json.loads(encoded)


def columns_fit(table, values):
    for name, value in values.items():
        column = table.c.get(name)
        if column is None:
            raise invalid()
        if value is None and not column.nullable:
            raise invalid()
        length = getattr(column.type, "length", None)
        if length and isinstance(value, str) and len(value) > length:
            raise invalid()


@dataclass(frozen=True, slots=True)
class ConfigImportPreparation:
    records: tuple[tuple[tuple[str, object], ...], ...] = field(repr=False)
    _issued: object = field(repr=False, compare=False)


def _parse_import(payload):
    try:
        validate_untrusted_xlsx_xml(payload)
        book = load_workbook(
            BytesIO(payload), read_only=True, data_only=True, keep_links=False
        )
    except (UnsafeXlsxError, InvalidFileException, BadZipFile, ValueError, OSError):
        raise invalid("CONFIG_IMPORT_INVALID_XLSX") from None
    records = []
    size = 0
    try:
        sheet = book.active
        if sheet is None:
            raise invalid("CONFIG_IMPORT_INVALID_XLSX")
        for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 1):
            if index > 1000:
                raise limit()
            if not row or len(row) < 2 or not row[1]:
                continue
            if len(row) < 7:
                raise invalid("CONFIG_IMPORT_INVALID_XLSX")
            values = {
                "config_name": str(row[0] or ""),
                "config_key": str(row[1]).strip(),
                "config_value": str(row[2] or ""),
                "config_type": str(row[3] or "text"),
                "config_group": str(row[4] or ""),
                "status": str(row[5] or "1"),
                "is_public": bool(row[6]) if row[6] is not None else False,
                "remark": str(row[7] or "") if len(row) > 7 and row[7] else None,
            }
            columns_fit(Config.__table__, values)
            if len(values["config_value"].encode("utf-8")) > 8192:
                raise limit()
            size += len(json.dumps(values, ensure_ascii=False).encode("utf-8"))
            if size > 1024 * 1024:
                raise limit()
            records.append(tuple(values.items()))
    finally:
        book.close()
    return ConfigImportPreparation(tuple(records), _IMPORT_ISSUED)


async def parse_import(payload: bytes) -> ConfigImportPreparation:
    if type(payload) is not bytes:
        raise invalid("CONFIG_IMPORT_INVALID_XLSX")
    if len(payload) > 10 * 1024 * 1024:
        raise limit()
    return await _import_worker.run(_parse_import, payload)


def freeze_input(operation, payload, target_id):
    if operation == "config.import":
        if (
            type(payload) is not ConfigImportPreparation
            or payload._issued is not _IMPORT_ISSUED
        ):
            raise invalid()
        return payload, None
    if operation.endswith(".batch_delete"):
        if type(payload) not in {tuple, list} or not payload:
            raise invalid()
        if len(payload) > 256:
            raise limit()
        return tuple(sorted({identifier(value) for value in payload})), None
    if operation.endswith(".delete"):
        return None, identifier(target_id)
    if isinstance(payload, BaseModel):
        payload = payload.model_dump(exclude_unset=operation.endswith(".update"))
    payload = bounded_json(payload, 1024 * 1024)
    if type(payload) is not dict:
        raise invalid()
    if operation == "settings.save":
        if (
            set(payload) != {"group", "values", "revision"}
            or payload["group"] not in SETTING_GROUPS
        ):
            raise invalid("SETTING_GROUP_INVALID")
        group, values = payload["group"], payload["values"]
        definitions = {key: d for key, d in SETTINGS.items() if d.group == group}
        if (
            type(values) is not dict
            or not values
            or not values.keys() <= definitions.keys()
            or type(payload["revision"]) is not str
        ):
            raise invalid("SETTING_VALUE_INVALID")
        normalized = {
            key: _validate(definitions[key], value) for key, value in values.items()
        }
        for value in normalized.values():
            if len(value.encode("utf-8")) > (
                400 * 1024 if group == "agreements" else 8192
            ):
                raise limit()
        return {**payload, "values": normalized}, None
    schema = {
        "menu.create": MenuCreate,
        "menu.update": MenuUpdate,
        "config.create": ConfigCreate,
        "config.update": ConfigUpdate,
        "mail.save": MailWrite,
    }.get(operation)
    if schema is None:
        raise invalid()
    try:
        model = schema.model_validate(payload)
    except (ValidationError, ValueError):
        raise invalid() from None
    values = model.model_dump(exclude_unset=operation.endswith(".update"))
    if operation.startswith("menu."):
        buttons = values.get("buttons")
        if buttons is not None:
            if len(buttons) > 128:
                raise limit()
            for button in buttons:
                columns_fit(
                    Menu.__table__,
                    {"permission": button["code"], "menu_name": button["desc"]},
                )
        if values.get("query") is not None:
            bounded_json(values["query"], 16 * 1024)
        if "parent_id" in values:
            values["parent_id"] = (
                None
                if values["parent_id"] in (None, 0)
                else identifier(values["parent_id"])
            )
        columns_fit(
            Menu.__table__,
            {k: v for k, v in values.items() if k not in {"buttons", "query"}},
        )
    if operation.startswith("config."):
        columns_fit(Config.__table__, values)
        if len((values.get("config_value") or "").encode("utf-8")) > 8192:
            raise limit()
    return values, identifier(target_id) if operation.endswith(".update") else None
