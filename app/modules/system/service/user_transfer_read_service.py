"""Explicit transfer metadata DTOs and scoped template dictionary sources."""

from datetime import UTC, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import func, select

from app.core.authorization_io import file_transform_worker
from app.core.exceptions import BusinessException, NotFoundException
from app.core.rbac import is_super_admin
from app.db.authorization_plan_session import _order_error
from app.modules.system.field_policy.evaluator import project_user_values
from app.modules.system.models.dept import Dept
from app.modules.system.models.role import Role
from app.modules.system.schemas.user_transfer import (
    UserExportTaskQuery,
    UserExportTaskResponse,
    UserImportBatchLogItem,
    UserImportBatchQuery,
    UserImportBatchResponse,
)
from app.modules.system.service import user_export_service as exports
from app.modules.system.service import user_import_service as imports
from app.modules.system.service import user_import_template_service as templates
from app.modules.system.service.authorization_control_identity import read_identity
from app.modules.system.service.grant_authority import grant_authority_service
from app.modules.system.service.user_field_context import load_current_user
from app.modules.system.service.user_import_validator import (
    _compute_accessible_dept_ids,
)
from app.modules.system.service.user_role_assignment_service import (
    user_role_assignment_service,
)
from app.modules.system.service.user_transfer_fields import user_transfer_fields


def batch_expires_at(batch):
    if batch is None or batch.created_at is None:
        return None
    if str(batch.status) in {"CREATED", "PREVIEW_DONE", "RUNNING"}:
        return batch.created_at + timedelta(minutes=10)
    return (batch.finished_at or batch.created_at) + timedelta(hours=24)


def batch_response(batch, operator_name):
    fields = {
        name: getattr(batch, name)
        for name in UserImportBatchResponse.model_fields
        if name not in {"operator_name", "sync_mode", "expires_at"}
    }
    fields.update(
        operator_name=operator_name, sync_mode=None, expires_at=batch_expires_at(batch)
    )
    return UserImportBatchResponse(**fields).model_dump(mode="json", by_alias=True)


def _build_template(depts, roles):
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    templates._build_data_sheet(ws)
    templates._build_instruction_sheet(wb.create_sheet("说明"))
    for name, headers, rows in (
        ("部门字典", templates._DEPT_DICT_COLUMNS, depts),
        ("角色字典", templates._ROLE_DICT_COLUMNS, roles),
    ):
        sheet = wb.create_sheet(name)
        sheet.append(
            [f"生成时间：{datetime.now(UTC).isoformat()}；当前授权范围，请使用最新模板"]
        )
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def _page(records, total, query):
    return {
        "records": records,
        "total": total,
        "current": query.current,
        "size": query.size,
    }


class UserTransferReadService:
    async def execute(
        self, db, tx, *, tenant, version, operation, query=None, identifier=None
    ):
        tx.assert_session(db)
        tx.assert_planning(tenant.tenant_id)
        if operation not in {
            "template",
            "batches",
            "batch",
            "logs",
            "exports",
            "export",
        }:
            raise _order_error()
        await read_identity(
            db,
            tenant=tenant,
            actor_user_version=version,
            permissions=("system:user:import",)
            if operation == "template"
            else ("system:user:list",),
        )
        context = await user_transfer_fields.context(
            db, tx, tenant=tenant, version=version
        )
        actor = await load_current_user(db, tenant=tenant, actor_user_version=version)
        cross_owner = is_super_admin(actor)
        if operation == "template":
            return await self.template(db, tx, tenant=tenant, actor=actor)
        if operation in {"exports", "batches"}:
            expected = (
                UserExportTaskQuery if operation == "exports" else UserImportBatchQuery
            )
            if type(query) is not expected:
                raise _order_error()
            query = expected.model_validate(query.model_dump())
        if operation == "exports":
            page = await exports.list_export_tasks(
                db,
                query,
                operator_id=tenant.actor_user_id,
                allow_cross_owner=cross_owner,
                tenant=tenant,
            )
            return _page(
                [self.export_dto(row) for row in page.records], page.total, query
            )
        if operation == "export":
            task = await exports.get_export_task(
                db,
                identifier,
                operator_id=tenant.actor_user_id,
                allow_cross_owner=cross_owner,
                tenant=tenant,
            )
            if task is None:
                raise NotFoundException(
                    "用户导出任务", error_code="AI_EXPORT_TASK_NOT_FOUND"
                )
            return self.export_dto(task)
        if operation == "batches":
            if not cross_owner:
                query.operator_id = tenant.actor_user_id
            rows, total = await imports.list_batches(db, query, tenant=tenant)
            names = await self.names(db, tx, context, rows)
            return _page(
                [self.batch_dto(row, names.get(row.operator_id)) for row, _ in rows],
                total,
                query,
            )
        batch, name = await imports.get_batch_detail(db, identifier, tenant=tenant)
        if (
            batch is None
            or not cross_owner
            and batch.operator_id != tenant.actor_user_id
        ):
            raise NotFoundException(
                "用户导入批次", error_code="AI_IMPORT_BATCH_NOT_FOUND"
            )
        if operation == "batch":
            names = await self.names(db, tx, context, ((batch, name),))
            return self.batch_dto(batch, names.get(batch.operator_id))
        if (
            type(query) is not dict
            or set(query) != {"event", "current", "size"}
            or type(query["current"]) is not int
            or query["current"] < 1
            or type(query["size"]) is not int
            or not 1 <= query["size"] <= 100
        ):
            raise _order_error()
        rows, total = await imports.list_batch_logs(
            db, identifier, tenant=tenant, **query
        )
        names = await self.names(db, tx, context, rows)
        records = []
        for log, _ in rows:
            detail = {
                key: value
                for key, value in log.detail.items()
                if key
                in {
                    "chunk_index",
                    "total_rows",
                    "chunk_size",
                    "failed_in_chunk",
                    "total_chunks",
                    "success_count",
                    "overwritten_count",
                    "skipped_count",
                    "failed_count",
                }
                and type(value) is int
                and value >= 0
            }
            records.append(
                UserImportBatchLogItem(
                    log_id=log.log_id,
                    event=log.event,
                    from_status=log.from_status,
                    to_status=log.to_status,
                    detail=detail,
                    operator_id=log.operator_id,
                    operator_name=names.get(log.operator_id),
                    created_at=log.created_at,
                ).model_dump(mode="json", by_alias=True)
            )
        return {
            "records": records,
            "total": total,
            "current": query["current"],
            "size": query["size"],
        }

    async def names(self, db, tx, context, rows):
        ids = tuple(sorted({row.operator_id for row, _ in rows}))
        decisions = await user_transfer_fields.decisions(
            db, tx, context, ids, operation="read", purpose="management.read"
        )
        return {
            row.operator_id: project_user_values(
                raw_values={"user_name": name},
                decisions=decisions[row.operator_id],
                entrypoint_caps={"user_name": "plain"},
            ).values.get("user_name")
            for row, name in rows
            if row.operator_id in decisions
        }

    @staticmethod
    def batch_dto(row, name):
        result = batch_response(row, name)
        result["filename"] = "users.xlsx"
        if row.failed_rows_file:
            result["failedRowsFile"] = f"/system/file/import-error/{row.batch_id}.xlsx"
        return result

    @staticmethod
    def export_dto(row):
        result = UserExportTaskResponse.model_validate(row).model_dump(
            mode="json", by_alias=True
        )
        result.update(filterSnapshot={}, reason="", errorMessage=None)
        return result

    async def template(self, db, tx, *, tenant, actor):
        tx.assert_planning(tenant.tenant_id)
        authority = await grant_authority_service.build(
            db, tenant.actor_user_id, tenant=tenant
        )
        departments = []
        if authority.allows_permission_codes({"system:dept:list"}):
            accessible = await _compute_accessible_dept_ids(db, actor, tenant=tenant)
            statement = select(Dept).where(
                Dept.tenant_id == tenant.tenant_id, Dept.status == "1"
            )
            if accessible is not None:
                statement = statement.where(Dept.dept_id.in_(accessible))
            departments = list(
                await db.scalars(statement.order_by(Dept.dept_id).limit(5001))
            )
        if len(departments) > 5000:
            raise BusinessException(
                422, "模板候选超出预算", error_code="AUTHORIZATION_PLAN_LIMIT_EXCEEDED"
            )
        roles = []
        if authority.allows_permission_codes({"system:user:role-auth"}):
            count = await db.scalar(
                select(func.pg_catalog.count(1))
                .select_from(Role)
                .where(Role.tenant_id == tenant.tenant_id)
            )
            if count > 1024:
                raise BusinessException(
                    422,
                    "模板候选超出预算",
                    error_code="AUTHORIZATION_PLAN_LIMIT_EXCEEDED",
                )
            roles = await user_role_assignment_service.list_assignable_roles(
                db,
                actor_user_id=tenant.actor_user_id,
                query=None,
                limit=1024,
                tenant=tenant,
            )
        lookup = {d.dept_id: d for d in departments}
        values = tuple(
            (
                d.dept_name,
                templates._build_dept_full_path(d, lookup),
                str(d.dept_id),
                d.status,
            )
            for d in departments
        )
        role_values = tuple((r.role_code, r.role_name, r.status) for r in roles)
        return await file_transform_worker.run(_build_template, values, role_values)


user_transfer_read_service = UserTransferReadService()
