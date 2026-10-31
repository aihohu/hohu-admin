"""Closed transfer dependencies, workbook projections and value-free provenance.

Manifests are evidence, never grants. The service rechecks current authority
under G/B before validating them or opening a stored file.
"""

import re
from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO

from openpyxl import Workbook

from app.core.exceptions import AuthorizationException
from app.modules.system.field_policy.documents import canonical_json
from app.modules.system.field_policy.evaluator import project_user_values
from app.modules.system.field_policy.query_plan import QueryDependency, UserQueryPlan
from app.modules.system.schemas.user_transfer import FailedRow

EXPORT_COLUMNS = (
    ("user_name", "user_name"),
    ("nickname", "nickname"),
    ("user_email", "user_email"),
    ("user_phone", "user_phone"),
    ("dept_id", "departments"),
    ("role_codes", "roles"),
    ("user_gender", "user_gender"),
    ("status", "status"),
    ("create_time", "create_time"),
)
_EXPORT_HEADERS = dict(
    zip(
        (name for name, _ in EXPORT_COLUMNS),
        (
            "账号",
            "昵称",
            "邮箱",
            "手机号",
            "部门",
            "角色编码",
            "性别",
            "状态",
            "创建时间",
        ),
        strict=True,
    )
)
_INPUT_ROOTS = {
    "user_name": "user_name",
    "employee_no": "employee_no",
    "nickname": "nickname",
    "user_email": "user_email",
    "user_phone": "user_phone",
    "user_gender": "user_gender",
    "status": "status",
    "dept_input": "departments",
    "role_input": "roles",
}


def transfer_denied():
    return AuthorizationException(
        "文件的当前字段授权无法确认，请重新生成",
        error_code="FIELD_TRANSFER_REGENERATE_REQUIRED",
    )


def export_query_plan(filters):
    definitions = {
        "user_names": ("user_name", "exact", "filter"),
        "user_name": ("user_name", "contains", "filter"),
        "nickname": ("nickname", "contains", "filter"),
        "user_email": ("user_email", "contains", "filter"),
        "user_phone": ("user_phone", "contains", "filter"),
        "status": ("status", "eq", "filter"),
        "dept_id": ("departments", "eq", "join"),
    }
    return UserQueryPlan(
        tuple(
            QueryDependency(*spec)
            for name, spec in definitions.items()
            if getattr(filters, name) is not None
        ),
        fixed_order=("user_id:asc",),
    )


def import_query_plan(*, employee_no):
    return UserQueryPlan(
        tuple(
            QueryDependency(name, "exact", "filter", "import")
            for name in (
                ("user_name", "employee_no") if employee_no else ("user_name",)
            )
        ),
        fixed_order=(),
    )


def import_write_plan(record, *, creating, rename):
    supplied = {
        _INPUT_ROOTS[name] for name in record.model_fields_set if name != "row_num"
    }
    # username is also the match selector; FULL_SYNC is its only overwrite path.
    scalar = supplied - {"user_name", "roles", "departments"}
    scalar |= {"user_gender", "status"}
    relations = {"departments"} | (
        {"roles"} if "roles" in supplied or creating else set()
    )
    if creating:
        scalar |= {
            "user_name",
            "password",
            "employee_no",
            "nickname",
            "user_email",
            "user_phone",
        }
    elif rename:
        scalar.add("user_name")
    return {
        "create" if creating else "update": frozenset(scalar),
        "replace": frozenset(relations),
    }


@dataclass(frozen=True, slots=True)
class ExportProjection:
    headers: tuple
    columns: tuple
    rows: tuple
    levels: tuple


def project_export(raw_rows, decisions):
    if len(raw_rows) != len(decisions) or len(raw_rows) > 5000:
        raise transfer_denied()
    projections = tuple(
        project_user_values(
            raw_values=raw,
            decisions=decision,
            entrypoint_caps={field: "plain" for _, field in EXPORT_COLUMNS},
        )
        for raw, decision in zip(raw_rows, decisions, strict=True)
    )
    columns = tuple(
        (name, field)
        for name, field in EXPORT_COLUMNS
        if any(p.field_access[field].allowed for p in projections)
    )
    return ExportProjection(
        tuple(_EXPORT_HEADERS[name] for name, _ in columns),
        tuple(field for _, field in columns),
        tuple(
            tuple(p.values.get(field, "未授权") for _, field in columns)
            for p in projections
        ),
        tuple(
            tuple(p.field_access[field].read for _, field in columns)
            for p in projections
        ),
    )


def build_export_workbook(projection):
    wb = Workbook()
    ws = wb.active
    ws.title = "users"
    for row in (projection.headers, *projection.rows):
        ws.append(list(row))
        # All values are data, even strings beginning with Excel formula syntax.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


_SAFE_CODES = frozenset(
    {
        "AI_IMPORT_USERNAME_INVALID",
        "AI_IMPORT_USERNAME_DUPLICATE",
        "AI_IMPORT_EMPLOYEE_NO_TOO_LONG",
        "AI_IMPORT_EMPLOYEE_NO_EXISTS",
        "AI_IMPORT_NICKNAME_TOO_LONG",
        "AI_IMPORT_EMAIL_INVALID",
        "AI_IMPORT_PHONE_INVALID",
        "AI_IMPORT_DEPT_INPUT_REQUIRED",
        "AI_IMPORT_GENDER_INVALID",
        "AI_IMPORT_STATUS_INVALID",
        "AI_IMPORT_BATCH_ABORTED",
        "AI_IMPORT_ADMIN_RENAME_FORBIDDEN",
        "AI_IMPORT_DEPT_NOT_FOUND",
        "AI_IMPORT_DEPT_AMBIGUOUS",
        "AI_IMPORT_ROLE_NOT_FOUND",
        "AI_IMPORT_ROLE_AMBIGUOUS",
        "AI_IMPORT_OUT_OF_SCOPE",
        "AI_IMPORT_DUPLICATE_TARGET",
        "USER_DEFAULT_ROLE_NOT_AVAILABLE",
        "AI_IMPORT_ROW_FAILED",
    }
)


def safe_failed_rows(rows):
    return [
        FailedRow(
            row_num=row.row_num,
            field=row.field if row.field in _INPUT_ROOTS else "_row",
            value="",
            reason="该行未通过导入校验，请检查输入和当前授权。",
            error_code=row.error_code
            if row.error_code in _SAFE_CODES
            else "AI_IMPORT_ROW_FAILED",
        )
        for row in rows
    ]


def make_lineage(
    *,
    tenant_id,
    owner_id,
    resource_id,
    kind,
    stamp,
    data,
    columns=(),
    sources=(),
    presence=(),
    requirements=(),
    mode=None,
):
    result = {
        "schema": "user-transfer-fields/1",
        "tenant_id": str(tenant_id),
        "owner_id": str(owner_id),
        "resource_id": resource_id,
        "kind": kind,
        "stamp": dict(stamp),
        "sha256": sha256(data).hexdigest(),
        "columns": list(columns),
        "sources": [[uid, list(levels)] for uid, levels in sources],
        "presence": [list(fields) for fields in presence],
        "requirements": list(requirements),
        "mode": mode,
    }
    validate_lineage(
        result,
        tenant_id=tenant_id,
        owner_id=owner_id,
        resource_id=resource_id,
        kind=kind,
        stamp=stamp,
    )
    return result


def validate_lineage(document, *, tenant_id, owner_id, resource_id, kind, stamp):
    keys = {
        "schema",
        "tenant_id",
        "owner_id",
        "resource_id",
        "kind",
        "stamp",
        "sha256",
        "columns",
        "sources",
        "presence",
        "requirements",
        "mode",
    }
    try:
        document_size = len(canonical_json(document).encode())
    except (TypeError, ValueError, OverflowError):
        raise transfer_denied() from None
    if (
        type(document) is not dict
        or document.keys() != keys
        or document_size > 2 * 1024 * 1024
        or document["schema"] != "user-transfer-fields/1"
        or document["tenant_id"] != str(tenant_id)
        or document["owner_id"] != str(owner_id)
        or document["resource_id"] != resource_id
        or document["kind"] != kind
        or kind not in {"export", "preview_source", "import_errors"}
        or document["stamp"] != stamp
        or not isinstance(document["sha256"], str)
        or not re.fullmatch(r"[0-9a-f]{64}", document["sha256"])
        or type(document["columns"]) is not list
        or type(document["sources"]) is not list
        or type(document["presence"]) is not list
        or type(document["requirements"]) is not list
        or len(document["requirements"]) > 2000
        or len(document["presence"]) > 2000
        or len(document["sources"]) > 5000
    ):
        raise transfer_denied()
    columns = document["columns"]
    if any(
        type(f) is not str or f not in {root for _, root in EXPORT_COLUMNS}
        for f in columns
    ) or len(columns) != len(set(columns)):
        raise transfer_denied()
    seen = set()
    for source in document["sources"]:
        if (
            type(source) is not list
            or len(source) != 2
            or type(source[0]) is not str
            or not re.fullmatch(r"[1-9][0-9]{0,18}", source[0])
            or int(source[0]) >= 2**63
            or source[0] in seen
            or type(source[1]) is not list
            or len(source[1]) != len(columns)
            or any(
                type(level) is not str or level not in {"plain", "masked", "hidden"}
                for level in source[1]
            )
        ):
            raise transfer_denied()
        seen.add(source[0])
    for fields in document["presence"]:
        if type(fields) is not list or any(
            type(f) is not str or f not in _INPUT_ROOTS for f in fields
        ):
            raise transfer_denied()
    if document["mode"] is not None and (
        type(document["mode"]) is not str
        or document["mode"] not in {"CREATE_ONLY", "UPDATE_PROFILE", "FULL_SYNC"}
    ):
        raise transfer_denied()
    for requirement in document["requirements"]:
        if type(requirement) is not dict or requirement.keys() != {
            "user_id",
            "dept_id",
            "operations",
        }:
            raise transfer_denied()
        for name in ("user_id", "dept_id"):
            value = requirement[name]
            if value is not None and (
                type(value) is not str
                or re.fullmatch(r"[1-9][0-9]{0,18}", value) is None
                or int(value) >= 2**63
            ):
                raise transfer_denied()
        operations = requirement["operations"]
        if type(operations) is not dict or set(operations) != {
            "create" if requirement["user_id"] is None else "update",
            "replace",
        }:
            raise transfer_denied()
        for fields in operations.values():
            if type(fields) is not list or any(
                type(f) is not str or f not in {*_INPUT_ROOTS.values(), "password"}
                for f in fields
            ):
                raise transfer_denied()
    if kind == "preview_source" and len(document["requirements"]) != len(
        document["presence"]
    ):
        raise transfer_denied()
    if type(document["stamp"]) is not dict or any(
        type(document["stamp"].get(k)) is not int or document["stamp"][k] < 1
        for k in ("global_epoch", "tenant_epoch")
    ):
        raise transfer_denied()
    return document


def verify_bytes(document, data):
    if sha256(data).hexdigest() != document["sha256"]:
        raise transfer_denied()
